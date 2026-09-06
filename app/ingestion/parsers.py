"""文档解析模块：把不同格式的文件提取成纯文本。

支持格式：
- .pdf  : 用 pypdf 逐页提取文本（仅支持"文本型 PDF"，扫描件需要 OCR，暂不支持）
- .docx : 用 python-docx 提取段落和表格文字
- .md / .txt : 直接读取文本
"""

from pathlib import Path

from docx import Document
from pypdf import PdfReader


class ParseError(Exception):
    """文档解析失败时抛出。错误信息会直接返回给前端用户。"""


def parse_file(path: str | Path) -> str:
    """根据扩展名分派到对应的解析器，返回整篇文档的纯文本。"""
    path = Path(path)
    ext = path.suffix.lower()

    if ext == ".pdf":
        text = _parse_pdf(path)
    elif ext == ".docx":
        text = _parse_docx(path)
    elif ext in (".md", ".txt"):
        text = _parse_text(path)
    else:
        raise ParseError(f"不支持的文件类型：{ext}（仅支持 .pdf / .docx / .md / .txt）")

    text = text.strip()
    if not text:
        raise ParseError(
            "未能从文件中提取到文本：如果是扫描件或图片版 PDF，需要 OCR 识别，本项目暂不支持"
        )
    return text


def _parse_pdf(path: Path) -> str:
    """逐页提取 PDF 文本。损坏或加密的文件转为友好错误。"""
    try:
        reader = PdfReader(str(path))
        pages: list[str] = []
        for page in reader.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                pages.append(page_text)
    except Exception as e:
        raise ParseError(f"PDF 解析失败：文件可能已损坏、加密或不是有效的 PDF（{e}）") from e
    return "\n".join(pages)


def _parse_docx(path: Path) -> str:
    """提取 Word 文档中的段落文字和表格内容。伪造/损坏的 docx 转为友好错误。"""
    try:
        doc = Document(str(path))
    except Exception as e:
        raise ParseError(f"Word 解析失败：文件可能已损坏或不是有效的 .docx（{e}）") from e

    parts: list[str] = []
    # 普通段落
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())

    # 表格：每行的单元格用 | 拼起来
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _parse_text(path: Path) -> str:
    """读取纯文本/Markdown。优先 UTF-8，失败则尝试 GBK（Windows 记事本默认编码）。"""
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return path.read_text(encoding="gbk", errors="ignore")
